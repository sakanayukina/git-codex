# -*- coding: utf-8 -*-
"""生成给导师的阶段汇报 DOCX：现状 + 数据 + 三条路线 + 待决策问题。"""
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ASSETS = r"C:\Users\18967\Desktop\git\assets"
OUT = os.path.join(ASSETS, "..", "震后救援仿真_阶段汇报.docx")
IMG_WIDE = os.path.join(ASSETS, "_v29_wide.png")
IMG_WALK = os.path.join(ASSETS, "_g1_walk.png")

CN_FONT = "微软雅黑"


def set_cn(run, size=None, bold=None, color=None):
    run.font.name = CN_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def para(doc, text, size=10.5, bold=False, space_after=6, color=None, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.25
    if align:
        p.alignment = align
    set_cn(p.add_run(text), size, bold, color)
    return p


def bullet(doc, text, size=10.5):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    set_cn(p.add_run(text), size)
    return p


def heading(doc, text, level=1):
    h = doc.add_heading(level=level)
    set_cn(h.add_run(text), 14 if level == 1 else 12, True, (0x1F, 0x38, 0x64))
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(6)
    return h


def table(doc, rows, widths=None, header=True, size=9.5):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            c = t.cell(i, j)
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            set_cn(p.add_run(str(cell)), size, bold=(header and i == 0))
    if widths:
        for j, w in enumerate(widths):
            for row in t.rows:
                row.cells[j].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2.0)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_cn(title.add_run("震后救援仿真环境与机器人运动验证"), 18, True, (0x1F, 0x38, 0x64))
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_cn(sub.add_run("阶段进展汇报与后续路线请示"), 11, False, (0x59, 0x59, 0x59))

    # 1 背景
    heading(doc, "一、工作背景与目标")
    para(doc, "目标是在 MuJoCo 中搭建一个可信的震后山地救援仿真环境，并把宇树 Go2（四足）"
              "与 G1（人形）导入其中，用于验证机器人在废墟地形中的可作业性，为后续任务规划与"
              "算法研究提供实验平台。")

    # 2 已完成
    heading(doc, "二、已完成的工作")
    bullet(doc, "真实地形：汶川 DEM，200 m × 200 m，落差 66.5 m，烘焙地面贴图，道路长在地形上。")
    bullet(doc, "村落场景：7 个可搜索点（H1 砖房废墟、H6 混凝土塌楼、H7 塌墙、H2 木屋等），"
                "含门洞碰撞、碎石带、堆积锥、208 棵树木，三角面 381 万。")
    bullet(doc, "机器人导入：Go2（URDF，16.3 kg / 20.8 万面）、G1（官方 MJCF，29 自由度与 12 自由度两版）。")
    bullet(doc, "验证场景：用宇树 terrain_tool 生成可控地形小院，三档难度（平地 / 台阶 / 台阶+碎石+起伏），"
                "用于逐级测试机器人通过性。")
    bullet(doc, "运动接入：G1 接上宇树官方预训练策略 motion.pt（12 自由度版），Go2 用手写步态。")

    # 3 数据
    heading(doc, "三、关键实验数据")
    para(doc, "（1）平地行走能力", bold=True, space_after=3)
    table(doc, [
        ["机器人", "控制方式", "8 秒位移", "机身高度稳定性", "航向漂移"],
        ["G1 人形", "官方 motion.pt", "3.64 m（约 0.45 m/s）", "0.77–0.80 m 全程稳定", "6°"],
        ["Go2 四足", "手写爬行步态", "2.58 m（6 秒）", "0.27–0.31 m", "12°"],
    ], widths=[3.0, 3.6, 4.2, 4.6, 2.4])

    para(doc, "（2）G1 抬腿能力曲线（同一策略，只改台阶高度）", bold=True, space_after=3)
    table(doc, [
        ["台阶高度", "8 秒位移", "结束骨盆高度", "结论"],
        ["0（平地）", "3.64 m", "0.77 m", "通过"],
        ["2 cm", "3.05 m", "0.90 m（踩上台阶）", "通过"],
        ["5 cm", "1.42 m", "0.13 m", "失败（摔倒）"],
        ["15 cm", "1.25 m", "0.13 m", "失败（摔倒）"],
    ], widths=[3.2, 3.4, 5.0, 5.0])
    para(doc, "结论：官方策略的抬腿上限约为 2–3 cm。而真实废墟场景中的台阶为 15 cm、瓦砾高 30 cm，"
              "远超其能力范围。", size=10, bold=True, color=(0xA3, 0x1D, 0x1D))

    para(doc, "（3）失效原因（已定位，非调参问题）", bold=True, space_after=3)
    bullet(doc, "宇树官方训练配置为 terrain.mesh_type = 'plane'，即在平地上训练。")
    bullet(doc, "训练与部署的本体必须完全一致：我们最初用 29 自由度模型裁剪（含手臂，35.1 kg）"
                "接入策略即摔倒；换成官方 12 自由度模型（无手臂，32.1 kg）后立刻可走。")

    # 4 配图
    heading(doc, "四、场景与效果")
    if os.path.exists(IMG_WIDE):
        doc.add_picture(IMG_WIDE, width=Cm(15.6))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cn(cap.add_run("图 1  震后村落仿真场景（H2 木屋、H1 砖房废墟、H6 混凝土塌楼）"),
               9, False, (0x59, 0x59, 0x59))
    if os.path.exists(IMG_WALK):
        doc.add_picture(IMG_WALK, width=Cm(7.2))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cn(cap.add_run("图 2  G1 使用官方预训练策略在平地上行走（骨盆高度稳定）"),
               9, False, (0x59, 0x59, 0x59))

    # 5 三条路线
    heading(doc, "五、后续三条技术路线")
    table(doc, [
        ["路线", "内容", "所需条件", "预计工时", "成果性质"],
        ["A", "接受现有能力边界，聚焦环境与任务规划研究", "无需额外算力", "1–2 周",
         "仿真平台 / 评测方法"],
        ["B", "本机自训策略（MuJoCo + PPO）", "本机 RTX 5060（8 GB）", "每轮 4–8 小时，可能多轮",
         "运动控制结果"],
        ["C", "在服务器上用 IsaacGym / IsaacLab 训练", "实验室或云 GPU，Linux 环境", "1–2 周（含环境搭建）",
         "与宇树训练配置完全对齐"],
    ], widths=[1.2, 5.4, 4.4, 3.6, 3.4])

    para(doc, "说明：IsaacGym 官方仅支持 Ubuntu，且其预编译版本对 RTX 5060 的 Blackwell 架构"
              "（sm_120）兼容性存疑，因此本机无法直接使用 IsaacGym，需借助服务器或云 GPU。",
         size=9.5, color=(0x59, 0x59, 0x59))

    # 6 待决策
    heading(doc, "六、需要请示的四个问题")
    table(doc, [
        ["序号", "问题", "影响"],
        ["1", "本课题定位是“构建仿真/评测平台”，还是“实现机器人在废墟中的运动”？",
         "前者现有成果已足够；后者必须自训策略"],
        ["2", "“官方策略在废墟地形失效”这一结论能否作为论文的一部分？",
         "可，则路线 A 即可成文且已有量化数据；不可，则需 B 或 C"],
        ["3", "是否有可用的服务器 / 云 GPU？是否要求与宇树训练配置严格对齐？",
         "有资源且要求严谨 → C；仅需结果 → B"],
        ["4", "时间预算是多少（多久需要产出）？", "紧 → A；有 2–4 周 → B 可行"],
    ], widths=[1.4, 8.0, 8.6])

    # 7 风险
    heading(doc, "七、风险与说明")
    bullet(doc, "路线 B 不保证一次成功：双足上台阶是公认难点，奖励函数需多轮调试。")
    bullet(doc, "8 GB 显存对当前模型规模足够（瓶颈在物理步进而非显存），但训练总时长不确定。")
    bullet(doc, "本机已安装 torch（CPU 版），切换至 CUDA 版约需 10 分钟、2.5 GB 下载量。")

    # 8 建议
    heading(doc, "八、我的建议")
    para(doc, "建议先明确课题定位。若目标是提供仿真平台与评测方法，现有成果（场景、机器人、"
              "通过性量化曲线）已经可以支撑阶段性产出；若目标是机器人运动控制，建议采用路线 B "
              "先用 Go2（四足容错高、1–2 小时可验证训练链路），再投入 G1 上台阶的训练。")

    # 9 附录
    heading(doc, "附录 A  运行环境")
    para(doc, "MuJoCo 3.3.2（Python 3.9）；NVIDIA RTX 5060 Laptop（8 GB，驱动 591.91）；"
              "PyTorch 2.8.0（当前 CPU 版，可切 CUDA 版）；工作目录 C:\\Users\\18967\\Desktop\\git。",
         size=9.0, space_after=4)

    heading(doc, "附录 B  关键文件与复现命令")
    table(doc, [
        ["用途", "文件 / 命令"],
        ["村落展示场景", "assets/demo_v29.xml（含两台机器人，静态摆位）"],
        ["验证小院（三档）", "assets/yard_{flat,stairs,full}.xml"],
        ["G1 可跑场景", "assets/props/robots/g1/g1_12dof_scene.xml"],
        ["官方策略", "assets/props/robots/g1/motion.pt（宇树预训练）"],
        ["平地行走", "python assets/_tools/g1_policy.py --viewer --vx 0.5"],
        ["台阶测试", "python assets/_tools/make_g1_yard.py --step-h 0.05，再跑 g1_policy.py"],
        ["Go2 步态", "python assets/_tools/run_walk.py"],
        ["场景自检", "python assets/_tools/check_yard.py"],
    ], widths=[4.6, 12.0], size=8.5)

    heading(doc, "附录 C  待解决问题清单")
    table(doc, [
        ["编号", "问题", "现状 / 已排查到哪一步"],
        ["P1", "机器人无法上废墟台阶（≥5 cm）",
         "策略训练地形为平地；需重训，非调参可解"],
        ["P2", "是否投入训练、用哪条路线",
         "待导师决策（见第六节四个问题）"],
        ["P3", "IsaacGym 本机不可用",
         "官方仅支持 Ubuntu；RTX 5060（Blackwell）兼容性存疑"],
        ["P4", "Go2 无官方预训练策略",
         "官方只有 G1/H1/H1-2 策略；Go2 为手写步态，长时不稳"],
        ["P5", "废墟场景台阶尺寸与机器人能力不匹配",
         "台阶 15 cm vs 上限 2–3 cm，需明确取舍"],
    ], widths=[1.6, 6.0, 9.0], size=8.5)

    heading(doc, "附录 D  两台机器人参数对比")
    table(doc, [
        ["项目", "Go2（四足）", "G1（29 自由度）", "G1（12 自由度，用于策略）"],
        ["整机质量", "16.3 kg", "35.1 kg", "32.1 kg"],
        ["自由度数", "19（含浮动基座）", "36（含浮动基座）", "19（含浮动基座）"],
        ["执行器", "12 个位置电机（kp=100）", "29 个位置电机（kp=150）",
         "12 个力矩电机（kp 100/150/40 分关节）"],
        ["三角面数", "20.8 万", "59.8 万", "去掉双臂网格"],
        ["站立高度", "机身 0.27 m", "骨盆 0.86 m（约 1.32 m 高）", "骨盆 0.80 m"],
        ["控制方式", "手写爬行步态", "固定基座（展示用）", "宇树官方 motion.pt"],
        ["平地行走", "2.58 m / 6 s", "—（不用于行走）", "3.64 m / 8 s"],
    ], widths=[3.0, 4.6, 4.6, 4.8], size=8.5)
    para(doc, "说明：G1 有两套模型。29 自由度版（含双臂）用于场景展示；12 自由度版（仅双腿，无手臂）"
              "与宇树预训练策略的训练本体一致，因此只有它能跑官方策略——这也解释了早期"
              "“同一策略在 35.1 kg 模型上必摔”的现象。", size=8.5, space_after=4)

    # 去掉末尾多余空段，避免多出一页
    while doc.paragraphs and not doc.paragraphs[-1].text.strip():
        el = doc.paragraphs[-1]._element
        el.getparent().remove(el)

    doc.save(OUT)
    print("saved:", OUT)


if __name__ == "__main__":
    main()
